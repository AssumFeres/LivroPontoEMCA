from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from banco import conectar, criar_banco
from config import ANO_PADRAO, CAMINHO_BANCO, PASTA_HORARIOS, TURMAS


SEPARADOR_PROFESSORES = re.compile(r"\s*/\s*")


def normalizar_texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()


def normalizar_sigla(valor) -> str:
    texto = normalizar_texto(valor).upper()
    if texto in {"", "0", "0.0", "NONE", "NAN"}:
        return ""
    return texto


def dividir_professores(nome_completo: str) -> list[str]:
    """
    A coluna E da aba Base pode trazer mais de um professor separados por '/'.
    Exemplo:
        RODNEY FAGUNDES DOS SANTOS / MANUEL DE LEMOS GASPAR

    Nesse caso a mesma aula é atribuída aos dois professores.
    """
    partes = [p.strip() for p in SEPARADOR_PROFESSORES.split(nome_completo or "")]

    resultado: list[str] = []
    vistos: set[str] = set()
    for parte in partes:
        if not parte:
            continue
        chave = parte.casefold()
        if chave not in vistos:
            vistos.add(chave)
            resultado.append(parte)
    return resultado


def localizar_aba(workbook, nome_desejado: str):
    """Localiza a aba ignorando espaços extras e diferença entre maiúsculas/minúsculas."""
    procurado = nome_desejado.strip().casefold()
    for nome_real in workbook.sheetnames:
        if nome_real.strip().casefold() == procurado:
            return workbook[nome_real]
    raise KeyError(
        f"Aba '{nome_desejado}' não encontrada. Abas disponíveis: {workbook.sheetnames}"
    )


def construir_dicionario_base(workbook) -> dict[str, list[str]]:
    """
    Retorna:
        {
            'MAT': ['MARCOS CESAR FARIA'],
            'OTB': ['RODNEY FAGUNDES DOS SANTOS', 'MANUEL DE LEMOS GASPAR'],
            ...
        }

    A aba Base usa:
        Coluna A = Código da Disciplina
        Coluna E = Nome Completo
    """
    ws = localizar_aba(workbook, "Base")
    dicionario: dict[str, list[str]] = {}

    for linha in ws.iter_rows(min_row=3, values_only=True):
        codigo = normalizar_sigla(linha[0] if len(linha) > 0 else None)
        nome = normalizar_texto(linha[4] if len(linha) > 4 else None)

        if not codigo or not nome:
            continue

        professores = dividir_professores(nome)
        if professores:
            dicionario[codigo] = professores

    return dicionario


def converter_data(valor) -> date | None:
    if valor is None:
        return None
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    return None


def registrar_log(conn, cfg: dict, tipo: str, mensagem: str) -> None:
    conn.execute(
        """
        INSERT INTO log_importacao (turma, arquivo, aba, tipo, mensagem)
        VALUES (?, ?, ?, ?, ?)
        """,
        (
            cfg["turma"],
            str(cfg["arquivo"]),
            cfg["aba_horario"],
            tipo,
            mensagem,
        ),
    )


def importar_turma(conn, cfg: dict, ano: int = ANO_PADRAO) -> int:
    caminho_arquivo = PASTA_HORARIOS / cfg["arquivo"]
    if not caminho_arquivo.exists():
        registrar_log(conn, cfg, "ERRO", f"Arquivo não encontrado: {caminho_arquivo}")
        print(f"[ERRO] Arquivo não encontrado: {caminho_arquivo}")
        return 0

    # data_only=True é indispensável: AULA 01 e AULA 02 são fórmulas nas planilhas.
    # Assim o Python lê o valor calculado e salvo pelo Excel, e não a fórmula HLOOKUP etc.
    wb = load_workbook(caminho_arquivo, data_only=True, read_only=True)

    try:
        base = construir_dicionario_base(wb)
        ws = localizar_aba(wb, cfg["aba_horario"])

        # Atualização segura: remove apenas os registros dessa turma no ano importado.
        conn.execute(
            "DELETE FROM aulas WHERE turma = ? AND substr(data, 1, 4) = ?",
            (cfg["turma"], str(ano)),
        )

        quantidade = 0

        for linha in ws.iter_rows(min_row=3, min_col=1, max_col=3, values_only=True):
            data_aula = converter_data(linha[0])
            if data_aula is None or data_aula.year != ano:
                continue

            sigla_01 = normalizar_sigla(linha[1])
            sigla_02 = normalizar_sigla(linha[2])

            # Dicionário temporário por professor para a mesma data/turma.
            # Assim o mesmo professor pode ter Aula 01 e Aula 02 na mesma linha.
            registros = defaultdict(lambda: {"aula_01": "", "aula_02": ""})

            if sigla_01:
                professores = base.get(sigla_01)
                if professores:
                    for professor in professores:
                        registros[professor]["aula_01"] = sigla_01
                else:
                    registrar_log(
                        conn,
                        cfg,
                        "SIGLA_NAO_ENCONTRADA",
                        f"{data_aula:%d/%m/%Y} - Aula 01: '{sigla_01}' não existe na aba Base.",
                    )

            if sigla_02:
                professores = base.get(sigla_02)
                if professores:
                    for professor in professores:
                        registros[professor]["aula_02"] = sigla_02
                else:
                    registrar_log(
                        conn,
                        cfg,
                        "SIGLA_NAO_ENCONTRADA",
                        f"{data_aula:%d/%m/%Y} - Aula 02: '{sigla_02}' não existe na aba Base.",
                    )

            for professor, aulas in registros.items():
                conn.execute(
                    """
                    INSERT INTO aulas (
                        turma, turno, data, nome_completo,
                        aula_01, aula_02, arquivo_origem, aba_origem
                    )
                    VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                    ON CONFLICT(turma, data, nome_completo)
                    DO UPDATE SET
                        turno = excluded.turno,
                        aula_01 = excluded.aula_01,
                        aula_02 = excluded.aula_02,
                        arquivo_origem = excluded.arquivo_origem,
                        aba_origem = excluded.aba_origem
                    """,
                    (
                        cfg["turma"],
                        cfg["turno"],
                        data_aula.isoformat(),
                        professor,
                        aulas["aula_01"],
                        aulas["aula_02"],
                        str(caminho_arquivo),
                        ws.title,
                    ),
                )
                quantidade += 1

        registrar_log(conn, cfg, "OK", f"{quantidade} registros importados.")
        return quantidade
    finally:
        wb.close()


def importar_todos(ano: int = ANO_PADRAO) -> int:
    criar_banco(CAMINHO_BANCO, ano)

    total = 0
    with conectar(CAMINHO_BANCO) as conn:
        for cfg in TURMAS:
            quantidade = importar_turma(conn, cfg, ano)
            total += quantidade
            print(f"{cfg['turma']}: {quantidade} registros")

    print(f"\nTotal importado: {total} registros")
    return total


if __name__ == "__main__":
    importar_todos()
