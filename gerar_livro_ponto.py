from __future__ import annotations

import calendar
import re
import shutil
import sqlite3
from copy import copy, deepcopy
from io import BytesIO
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage

from config import CAMINHO_BANCO, MESES_PT, PASTA_EXPORTACOES, PASTA_PROJETO


PASTA_MODELOS = PASTA_PROJETO / "modelos"
CAMINHO_MODELO_PONTO = PASTA_MODELOS / "Modelo_de_Ponto_Mensal_Instrutor.xlsx"

ABA_MODELO = "Prof. (2)"

LINHA_INICIAL_DIAS = 9
LINHA_FINAL_DIAS = 31

COLUNAS_VESPERTINO = {
    "aula_01": ("C", "D", "E"),
    "aula_02": ("F", "G", "H"),
}

COLUNAS_NOTURNO = {
    "aula_01": ("I", "J", "K"),
    "aula_02": ("L", "M", "N"),
}

DIAS_SEMANA_PT = {
    0: "Segunda-feira",
    1: "Terça-feira",
    2: "Quarta-feira",
    3: "Quinta-feira",
    4: "Sexta-feira",
    5: "Sábado",
    6: "Domingo",
}


def dias_uteis_mes(ano: int, mes: int) -> list[date]:
    ultimo_dia = calendar.monthrange(ano, mes)[1]
    return [
        date(ano, mes, dia)
        for dia in range(1, ultimo_dia + 1)
        if date(ano, mes, dia).weekday() <= 4
    ]


def normalizar_turno(turno: str | None, turma: str) -> str:
    if turno:
        turno_normalizado = turno.strip().upper()
        if turno_normalizado in {"VESPERTINO", "NOTURNO"}:
            return turno_normalizado
    return "VESPERTINO" if turma.strip().upper() == "U3" else "NOTURNO"


def sigla_valida(valor: str | None) -> str | None:
    if valor is None:
        return None
    texto = str(valor).strip()
    return texto or None


def juntar_siglas(siglas: list[str]) -> str | None:
    unicas: list[str] = []
    for sigla in siglas:
        sigla = sigla.strip()
        if sigla and sigla not in unicas:
            unicas.append(sigla)
    return " / ".join(unicas) if unicas else None


def preencher_tres_colunas(ws, linha: int, colunas: tuple[str, str, str], valor: str | None) -> None:
    for coluna in colunas:
        ws[f"{coluna}{linha}"] = valor


def limpar_area_dias(ws) -> None:
    for linha in range(LINHA_INICIAL_DIAS, LINHA_FINAL_DIAS + 1):
        for coluna in range(1, 17):
            ws.cell(row=linha, column=coluna).value = None


def obter_colunas_tabela(conexao: sqlite3.Connection, tabela: str) -> set[str]:
    linhas = conexao.execute(f"PRAGMA table_info({tabela})").fetchall()
    return {str(linha[1]) for linha in linhas}


def obter_coluna_professor(conexao: sqlite3.Connection) -> str:
    colunas = obter_colunas_tabela(conexao, "aulas")
    for candidato in ("professor", "nome_completo", "nome_professor", "instrutor"):
        if candidato in colunas:
            return candidato
    raise RuntimeError(
        "Não foi encontrada na tabela 'aulas' uma coluna de professor. "
        f"Colunas existentes: {', '.join(sorted(colunas))}"
    )


def obter_coluna_turno(conexao: sqlite3.Connection) -> str | None:
    colunas = obter_colunas_tabela(conexao, "aulas")
    return "turno" if "turno" in colunas else None


def obter_professores_mes(conexao: sqlite3.Connection, ano: int, mes: int) -> list[str]:
    inicio = f"{ano:04d}-{mes:02d}-01"
    proximo = f"{ano + 1:04d}-01-01" if mes == 12 else f"{ano:04d}-{mes + 1:02d}-01"

    coluna_professor = obter_coluna_professor(conexao)

    sql = f'''
        SELECT DISTINCT "{coluna_professor}"
        FROM aulas
        WHERE data >= ?
          AND data < ?
          AND "{coluna_professor}" IS NOT NULL
          AND TRIM("{coluna_professor}") <> ''
        ORDER BY "{coluna_professor}"
    '''

    linhas = conexao.execute(sql, (inicio, proximo)).fetchall()
    return [str(linha[0]).strip() for linha in linhas]


def obter_aulas_professor(
    conexao: sqlite3.Connection,
    professor: str,
    ano: int,
    mes: int,
) -> list[sqlite3.Row]:
    inicio = f"{ano:04d}-{mes:02d}-01"
    proximo = f"{ano + 1:04d}-01-01" if mes == 12 else f"{ano:04d}-{mes + 1:02d}-01"

    coluna_professor = obter_coluna_professor(conexao)
    coluna_turno = obter_coluna_turno(conexao)
    expressao_turno = '"turno"' if coluna_turno else "NULL"

    sql = f'''
        SELECT
            turma,
            {expressao_turno} AS turno,
            data,
            "{coluna_professor}" AS professor,
            aula_01,
            aula_02
        FROM aulas
        WHERE "{coluna_professor}" = ?
          AND data >= ?
          AND data < ?
        ORDER BY data, turma
    '''

    return conexao.execute(sql, (professor, inicio, proximo)).fetchall()


def obter_resumo_professor(aulas: list[sqlite3.Row]) -> tuple[str, str]:
    disciplinas: set[str] = set()
    turmas: set[str] = set()

    for aula in aulas:
        turmas.add(str(aula["turma"]).strip())

        aula_01 = sigla_valida(aula["aula_01"])
        aula_02 = sigla_valida(aula["aula_02"])

        if aula_01:
            disciplinas.add(aula_01)
        if aula_02:
            disciplinas.add(aula_02)

    return ", ".join(sorted(disciplinas)), ", ".join(sorted(turmas))


def agrupar_aulas_por_data(aulas: list[sqlite3.Row]) -> dict[str, dict[str, dict[str, list[str]]]]:
    dados = defaultdict(
        lambda: {
            "VESPERTINO": {"aula_01": [], "aula_02": []},
            "NOTURNO": {"aula_01": [], "aula_02": []},
        }
    )

    for aula in aulas:
        data_aula = str(aula["data"])
        turma = str(aula["turma"])
        turno = normalizar_turno(aula["turno"], turma)

        aula_01 = sigla_valida(aula["aula_01"])
        aula_02 = sigla_valida(aula["aula_02"])

        if aula_01:
            dados[data_aula][turno]["aula_01"].append(aula_01)
        if aula_02:
            dados[data_aula][turno]["aula_02"].append(aula_02)

    return dict(dados)


def nome_aba_professor(nome: str, existentes: set[str]) -> str:
    nome = re.sub(r'[\\/*?:\[\]]', "", nome.strip())
    nome = re.sub(r"\s+", " ", nome)
    nome = nome[:31] if len(nome) > 31 else nome

    base = nome or "Professor"
    candidato = base
    contador = 2

    while candidato in existentes:
        sufixo = f"_{contador}"
        candidato = f"{base[:31-len(sufixo)]}{sufixo}"
        contador += 1

    return candidato


def capturar_imagens_modelo(ws_modelo) -> list[dict]:
    """
    Lê cada imagem da aba modelo UMA ÚNICA VEZ e guarda seus bytes,
    dimensões e posição.

    Isso evita o erro:
        I/O operation on closed file

    O openpyxl fecha o fluxo da imagem original após a primeira leitura.
    """
    imagens_cache: list[dict] = []

    for imagem in getattr(ws_modelo, "_images", []):
        dados = imagem._data()

        imagens_cache.append(
            {
                "dados": dados,
                "width": imagem.width,
                "height": imagem.height,
                "anchor": deepcopy(imagem.anchor),
            }
        )

    return imagens_cache


def aplicar_imagens_na_planilha(ws_destino, imagens_cache: list[dict]) -> None:
    """
    Cria uma nova imagem independente em cada aba de professor.
    """
    for item in imagens_cache:
        fluxo = BytesIO(item["dados"])

        nova_imagem = ExcelImage(fluxo)
        nova_imagem.width = item["width"]
        nova_imagem.height = item["height"]
        nova_imagem.anchor = deepcopy(item["anchor"])

        ws_destino.add_image(nova_imagem)


def preencher_planilha_professor(
    ws,
    professor: str,
    ano: int,
    mes: int,
    aulas: list[sqlite3.Row],
) -> None:
    disciplinas, turmas = obter_resumo_professor(aulas)

    ws["B3"] = professor
    ws["B4"] = datetime(ano, mes, 1)
    ws["B4"].number_format = "mmmm/yyyy"
    ws["B5"] = disciplinas
    ws["O5"] = turmas

    limpar_area_dias(ws)

    dias = dias_uteis_mes(ano, mes)
    capacidade = LINHA_FINAL_DIAS - LINHA_INICIAL_DIAS + 1

    if len(dias) > capacidade:
        raise ValueError(
            f"O modelo possui espaço para {capacidade} dias úteis, "
            f"mas {MESES_PT[mes]}/{ano} possui {len(dias)}."
        )

    aulas_por_data = agrupar_aulas_por_data(aulas)

    for indice, data_atual in enumerate(dias):
        linha = LINHA_INICIAL_DIAS + indice

        ws[f"A{linha}"] = datetime(data_atual.year, data_atual.month, data_atual.day)
        ws[f"A{linha}"].number_format = "dd/mm/yyyy"
        ws[f"B{linha}"] = DIAS_SEMANA_PT[data_atual.weekday()]

        dados_dia = aulas_por_data.get(data_atual.isoformat())
        if not dados_dia:
            continue

        vespertino_01 = juntar_siglas(dados_dia["VESPERTINO"]["aula_01"])
        vespertino_02 = juntar_siglas(dados_dia["VESPERTINO"]["aula_02"])
        noturno_01 = juntar_siglas(dados_dia["NOTURNO"]["aula_01"])
        noturno_02 = juntar_siglas(dados_dia["NOTURNO"]["aula_02"])

        preencher_tres_colunas(ws, linha, COLUNAS_VESPERTINO["aula_01"], vespertino_01)
        preencher_tres_colunas(ws, linha, COLUNAS_VESPERTINO["aula_02"], vespertino_02)
        preencher_tres_colunas(ws, linha, COLUNAS_NOTURNO["aula_01"], noturno_01)
        preencher_tres_colunas(ws, linha, COLUNAS_NOTURNO["aula_02"], noturno_02)

        if any((vespertino_01, vespertino_02, noturno_01, noturno_02)):
            ws[f"O{linha}"] = "X"

    for linha in range(LINHA_INICIAL_DIAS, LINHA_FINAL_DIAS + 1):
        ws[f"P{linha}"] = None

    ws.sheet_view.showGridLines = False
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def gerar_livros_ponto(
    ano: int,
    mes: int,
    caminho_banco: str | Path = CAMINHO_BANCO,
    caminho_modelo: str | Path = CAMINHO_MODELO_PONTO,
) -> list[Path]:
    if not 1 <= mes <= 12:
        raise ValueError("Mês deve estar entre 1 e 12.")

    caminho_banco = Path(caminho_banco)
    caminho_modelo = Path(caminho_modelo)

    if not caminho_banco.exists():
        raise FileNotFoundError(f"Banco de dados não encontrado: {caminho_banco}")

    if not caminho_modelo.exists():
        raise FileNotFoundError(f"Modelo não encontrado: {caminho_modelo}")

    pasta_saida = PASTA_EXPORTACOES / "livros_ponto"
    pasta_saida.mkdir(parents=True, exist_ok=True)

    caminho_saida = pasta_saida / f"LIVRO_PONTO_{ano}_{mes:02d}_{MESES_PT[mes]}.xlsx"

    shutil.copy2(caminho_modelo, caminho_saida)

    conexao = sqlite3.connect(caminho_banco)
    conexao.row_factory = sqlite3.Row

    try:
        professores = obter_professores_mes(conexao, ano, mes)

        if not professores:
            raise ValueError(
                f"Nenhum professor com aula encontrado em {MESES_PT[mes]}/{ano}."
            )

        wb = load_workbook(caminho_saida)

        if ABA_MODELO not in wb.sheetnames:
            raise ValueError(f"A aba '{ABA_MODELO}' não foi encontrada no modelo.")

        ws_modelo = wb[ABA_MODELO]

        # Lê as imagens do modelo apenas uma vez.
        imagens_cache = capturar_imagens_modelo(ws_modelo)

        for nome in list(wb.sheetnames):
            if nome != ABA_MODELO:
                del wb[nome]

        existentes = set(wb.sheetnames)

        print(f"\nGerando livro ponto consolidado de {MESES_PT[mes]}/{ano}...")
        print(f"Professores encontrados: {len(professores)}\n")

        for numero, professor in enumerate(professores, start=1):
            aulas = obter_aulas_professor(conexao, professor, ano, mes)

            ws = wb.copy_worksheet(ws_modelo)
            aplicar_imagens_na_planilha(ws, imagens_cache)

            nome_aba = nome_aba_professor(professor, existentes)
            ws.title = nome_aba
            existentes.add(nome_aba)

            preencher_planilha_professor(
                ws=ws,
                professor=professor,
                ano=ano,
                mes=mes,
                aulas=aulas,
            )

            print(f"[{numero:02d}/{len(professores):02d}] {professor}")

        del wb[ABA_MODELO]
        wb.active = 0
        wb.save(caminho_saida)

        print("\nArquivo gerado com sucesso em:")
        print(caminho_saida)

        return [caminho_saida]

    finally:
        conexao.close()


if __name__ == "__main__":
    ano = int(input("Ano: ").strip())
    mes = int(input("Mês (1-12): ").strip())

    arquivos = gerar_livros_ponto(ano=ano, mes=mes)

    print(f"\nTotal de arquivos gerados: {len(arquivos)}")
