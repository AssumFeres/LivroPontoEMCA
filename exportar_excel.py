from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from banco import conectar
from config import CAMINHO_BANCO, MESES_PT, PASTA_EXPORTACOES
from resumo_mensal import atualizar_resumo_mensal


COR_CABECALHO = "002060"
COR_TEXTO_CABECALHO = "FFFFFF"
BORDA = Side(style="thin", color="B7B7B7")


def ajustar_larguras(ws, limites: dict[int, int] | None = None) -> None:
    limites = limites or {}
    for coluna in range(1, ws.max_column + 1):
        maior = 0
        for linha in range(1, ws.max_row + 1):
            valor = ws.cell(linha, coluna).value
            if valor is not None:
                maior = max(maior, len(str(valor)))
        largura = min(max(maior + 2, 10), limites.get(coluna, 40))
        ws.column_dimensions[get_column_letter(coluna)].width = largura


def estilizar_cabecalho(ws) -> None:
    for celula in ws[1]:
        celula.fill = PatternFill("solid", fgColor=COR_CABECALHO)
        celula.font = Font(color=COR_TEXTO_CABECALHO, bold=True)
        celula.alignment = Alignment(horizontal="center", vertical="center")
        celula.border = Border(top=BORDA, bottom=BORDA, left=BORDA, right=BORDA)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def exportar_mes(ano: int, mes: int, caminho_saida: str | Path | None = None) -> Path:
    if mes not in range(1, 13):
        raise ValueError("O mês deve estar entre 1 e 12.")

    atualizar_resumo_mensal(ano, mes)

    inicio = f"{ano:04d}-{mes:02d}-01"
    if mes == 12:
        fim = f"{ano + 1:04d}-01-01"
    else:
        fim = f"{ano:04d}-{mes + 1:02d}-01"

    with conectar(CAMINHO_BANCO) as conn:
        aulas = conn.execute(
            """
            SELECT turma, data, nome_completo, aula_01, aula_02
            FROM aulas
            WHERE data >= ? AND data < ?
            ORDER BY data,
                     CASE turno WHEN 'VESPERTINO' THEN 1 ELSE 2 END,
                     turma,
                     nome_completo
            """,
            (inicio, fim),
        ).fetchall()

        resumo = conn.execute(
            """
            SELECT nome_completo, disciplinas, turmas
            FROM resumo_professor_mensal
            WHERE ano = ? AND mes = ?
            ORDER BY nome_completo
            """,
            (ano, mes),
        ).fetchall()

    if caminho_saida is None:
        PASTA_EXPORTACOES.mkdir(parents=True, exist_ok=True)
        caminho_saida = PASTA_EXPORTACOES / (
            f"LIVRO_PONTO_{ano}_{mes:02d}_{MESES_PT[mes]}.xlsx"
        )
    else:
        caminho_saida = Path(caminho_saida)
        caminho_saida.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # ---------------------------------------------------------
    # ABA 1 - LIVRO PONTO
    # ---------------------------------------------------------
    ws = wb.active
    ws.title = "Livro Ponto"
    ws.append(["TURMA", "DATA", "NOME COMPLETO", "AULA 01", "AULA 02"])

    for turma, data_iso, professor, aula_01, aula_02 in aulas:
        ano_l, mes_l, dia_l = map(int, data_iso.split("-"))
        from datetime import date

        ws.append([
            turma,
            date(ano_l, mes_l, dia_l),
            professor,
            aula_01 or "",
            aula_02 or "",
        ])

    for celula in ws["B"][1:]:
        celula.number_format = "dd/mm/yyyy"

    estilizar_cabecalho(ws)
    ajustar_larguras(ws, {1: 14, 2: 14, 3: 45, 4: 16, 5: 16})

    for linha in ws.iter_rows(min_row=2):
        for celula in linha:
            celula.border = Border(top=BORDA, bottom=BORDA, left=BORDA, right=BORDA)
            celula.alignment = Alignment(vertical="center")

    # ---------------------------------------------------------
    # ABA 2 - RESUMO PROFESSORES
    # ---------------------------------------------------------
    ws2 = wb.create_sheet("Resumo Professores")
    ws2.append(["NOME DO PROFESSOR", "DISCIPLINAS DO MÊS", "TURMAS DO MÊS"])

    for professor, disciplinas, turmas in resumo:
        ws2.append([professor, disciplinas, turmas])

    estilizar_cabecalho(ws2)
    ajustar_larguras(ws2, {1: 45, 2: 50, 3: 35})

    for linha in ws2.iter_rows(min_row=2):
        for celula in linha:
            celula.border = Border(top=BORDA, bottom=BORDA, left=BORDA, right=BORDA)
            celula.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(caminho_saida)
    return caminho_saida


if __name__ == "__main__":
    ano = int(input("Ano: "))
    mes = int(input("Mês (1-12): "))
    arquivo = exportar_mes(ano, mes)
    print(f"Arquivo criado: {arquivo}")
